#!/usr/bin/env python3
"""
Convert Markdown to Word document for software copyright application.
Converts 软件设计说明书.md to 软件设计说明书.docx with proper Chinese formatting.

Usage:
    python md_to_word.py

Dependencies:
    pip install python-docx pdf2image pillow

Note: For PDF to PNG conversion, poppler-utils must be installed:
    sudo apt-get install poppler-utils
"""

import re
import os
import tempfile
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pdf2image import convert_from_path


def set_chinese_font(run, font_name='宋体', font_size=12, bold=False):
    """Set Chinese font for a run."""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.insert(0, rFonts)


def add_page_number(doc):
    """Add page numbers to document footer."""
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        run = p.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        
        instrText = OxmlElement('w:instrText')
        instrText.text = "PAGE"
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)


def convert_pdf_to_png(pdf_path, output_dir):
    """Convert PDF to PNG using pdf2image."""
    try:
        images = convert_from_path(pdf_path, dpi=200)
        if images:
            output_path = os.path.join(output_dir, Path(pdf_path).stem + '.png')
            images[0].save(output_path, 'PNG')
            return output_path
    except Exception as e:
        print(f"Error converting PDF {pdf_path}: {e}")
    return None


def find_mermaid_image(mermaid_dir, counter, base_name="软件设计说明书"):
    """
    Find the pre-rendered mermaid image file for a given counter.
    
    The images are named like: 软件设计说明书_graph_01.png, 软件设计说明书_flowchart_02.png, etc.
    
    Args:
        mermaid_dir: Directory containing rendered Mermaid diagrams
        counter: The sequential counter of the mermaid diagram (1-based)
        base_name: Base name prefix for the image files
        
    Returns:
        Path to the image file if found, None otherwise
    """
    if not mermaid_dir or not os.path.isdir(mermaid_dir):
        return None
    
    # Format counter as two-digit string
    counter_str = f"{counter:02d}"
    
    # Search for image files matching the pattern: base_name_*_counter.png
    for filename in os.listdir(mermaid_dir):
        if filename.startswith(base_name) and filename.endswith(f"_{counter_str}.png"):
            return os.path.join(mermaid_dir, filename)
    
    return None


def get_image_dimensions_for_word(image_path, max_width_inches=5.5, max_height_inches=7.0):
    """
    Calculate appropriate width and height for an image in a Word document.
    
    Ensures the image fits within the specified maximum dimensions while
    maintaining aspect ratio.
    
    Args:
        image_path: Path to the image file
        max_width_inches: Maximum width in inches (default 5.5 for standard margins)
        max_height_inches: Maximum height in inches (default 7.0 to fit on one page)
        
    Returns:
        Tuple of (width_inches, height_inches) or (max_width_inches, None) if unable to read image
    """
    try:
        from PIL import Image
        with Image.open(image_path) as img:
            width_px, height_px = img.size
            aspect_ratio = width_px / height_px
            
            # Calculate dimensions if constrained by width
            width_by_width = max_width_inches
            height_by_width = max_width_inches / aspect_ratio
            
            # Calculate dimensions if constrained by height
            height_by_height = max_height_inches
            width_by_height = max_height_inches * aspect_ratio
            
            # Choose the smaller set of dimensions
            if height_by_width <= max_height_inches:
                return (width_by_width, height_by_width)
            else:
                return (width_by_height, height_by_height)
    except Exception as e:
        print(f"Warning: Could not read image dimensions for {image_path}: {e}")
        return (max_width_inches, None)


def convert_md_to_word(md_file, output_file, title="软件设计说明书", base_dir=None, mermaid_dir=None):
    """
    Convert Markdown to Word with improved formatting.
    
    Args:
        md_file: Path to input Markdown file
        output_file: Path to output Word document
        title: Document title for header
        base_dir: Base directory for resolving relative image paths
        mermaid_dir: Directory containing rendered Mermaid diagrams (optional)
    """
    
    if base_dir is None:
        base_dir = os.path.dirname(md_file)
    
    temp_dir = tempfile.mkdtemp()
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    doc = Document()
    
    # Set default style
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # Set page margins and header
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        
        header = section.header
        header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        header_run = header_para.add_run(f"{title}")
        set_chinese_font(header_run, '宋体', 9)
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    add_page_number(doc)
    
    lines = content.split('\n')
    in_code_block = False
    code_content = []
    code_lang = ""
    mermaid_counter = 0
    
    # Track context for merging numbered lists
    in_dev_goals = False
    in_tech_features = False
    pending_items = []
    
    def flush_pending():
        nonlocal pending_items
        if pending_items:
            p = doc.add_paragraph()
            text = "；".join(pending_items) + "。"
            run = p.add_run(text)
            set_chinese_font(run, '宋体', 10.5)
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.line_spacing = 1.5
            pending_items = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        
        # Handle code blocks
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip()
                code_content = []
            else:
                in_code_block = False
                if code_lang.lower() == 'mermaid':
                    mermaid_counter += 1
                    # Check if rendered mermaid image exists
                    mermaid_img_path = find_mermaid_image(mermaid_dir, mermaid_counter)
                    
                    if mermaid_img_path:
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run()
                        # Calculate appropriate size for the image
                        img_width, img_height = get_image_dimensions_for_word(mermaid_img_path)
                        if img_height:
                            run.add_picture(mermaid_img_path, width=Inches(img_width), height=Inches(img_height))
                        else:
                            run.add_picture(mermaid_img_path, width=Inches(img_width))
                    else:
                        # Placeholder for mermaid diagram
                        p = doc.add_paragraph()
                        run = p.add_run('[Mermaid图表 - 请参见原始Markdown文件]')
                        set_chinese_font(run, '楷体', 10)
                        run.italic = True
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    # Regular code block
                    for code_line in code_content:
                        p = doc.add_paragraph()
                        run = p.add_run(code_line)
                        run.font.name = 'Consolas'
                        run.font.size = Pt(9)
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                code_content = []
                code_lang = ""
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # Handle headers
        if line.startswith('# '):
            flush_pending()
            in_dev_goals = False
            in_tech_features = False
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            set_chinese_font(run, '黑体', 16, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
        elif line.startswith('## '):
            flush_pending()
            in_dev_goals = False
            in_tech_features = False
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            set_chinese_font(run, '黑体', 14, bold=True)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif line.startswith('### '):
            flush_pending()
            in_dev_goals = False
            in_tech_features = False
            p = doc.add_paragraph()
            run = p.add_run(line[4:])
            set_chinese_font(run, '黑体', 12, bold=True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif line.startswith('#### '):
            flush_pending()
            p = doc.add_paragraph()
            run = p.add_run(line[5:])
            set_chinese_font(run, '黑体', 11, bold=True)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
        
        # Handle horizontal rules
        elif line.strip() == '---':
            flush_pending()
            in_dev_goals = False
            in_tech_features = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        
        # Handle images
        elif line.strip().startswith('!['):
            flush_pending()
            match = re.match(r'!\[(.*?)\]\((.*?)\)', line.strip())
            if match:
                alt_text = match.group(1)
                img_path = match.group(2)
                full_img_path = os.path.join(base_dir, img_path)
                
                if img_path.lower().endswith('.pdf') and os.path.exists(full_img_path):
                    png_path = convert_pdf_to_png(full_img_path, temp_dir)
                    if png_path and os.path.exists(png_path):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run()
                        # Calculate appropriate size for the image
                        img_width, img_height = get_image_dimensions_for_word(png_path)
                        if img_height:
                            run.add_picture(png_path, width=Inches(img_width), height=Inches(img_height))
                        else:
                            run.add_picture(png_path, width=Inches(img_width))
                    else:
                        p = doc.add_paragraph()
                        run = p.add_run(f'[图片: {alt_text}]')
                        set_chinese_font(run, '楷体', 10)
                        run.italic = True
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(f'[图片: {alt_text}]')
                    set_chinese_font(run, '楷体', 10)
                    run.italic = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Handle bold figure captions
        elif line.strip().startswith('**图'):
            flush_pending()
            p = doc.add_paragraph()
            text = line.strip().replace('**', '')
            run = p.add_run(text)
            set_chinese_font(run, '宋体', 10, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
        
        # Handle bold subtitles like **开发目标：** or **技术特点：**
        elif re.match(r'^\*\*(.+?[：:])?\*\*\s*$', line.strip()):
            flush_pending()
            p = doc.add_paragraph()
            text = line.strip().replace('**', '')
            run = p.add_run(text)
            set_chinese_font(run, '黑体', 10.5, bold=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            
            # Track context
            if '开发目标' in text:
                in_dev_goals = True
                in_tech_features = False
            elif '技术特点' in text:
                in_dev_goals = False
                in_tech_features = True
            else:
                in_dev_goals = False
                in_tech_features = False
        
        # Handle bullet points
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            flush_pending()
            text = line.strip()[2:]
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(text)
            set_chinese_font(run, '宋体', 10.5)
        
        # Handle numbered lists
        elif re.match(r'^\d+\.\s', line.strip()):
            text = re.sub(r'^\d+\.\s', '', line.strip())
            
            # Merge numbered items for 开发目标 section
            if in_dev_goals and not text.startswith('**'):
                pending_items.append(text)
            else:
                flush_pending()
                # Handle bold prefix in tech features
                bold_match = re.match(r'\*\*(.+?)\*\*[：:]\s*(.+)', text)
                if bold_match:
                    p = doc.add_paragraph()
                    p.paragraph_format.first_line_indent = Pt(21)
                    run1 = p.add_run(bold_match.group(1) + "：")
                    set_chinese_font(run1, '黑体', 10.5, bold=True)
                    run2 = p.add_run(bold_match.group(2))
                    set_chinese_font(run2, '宋体', 10.5)
                else:
                    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
                    p = doc.add_paragraph(style='List Number')
                    run = p.add_run(text)
                    set_chinese_font(run, '宋体', 10.5)
        
        # Handle tables
        elif line.strip().startswith('|'):
            flush_pending()
            table_lines = [line]
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith('|'):
                table_lines.append(lines[j])
                j += 1
            
            data_lines = [l for l in table_lines if not re.match(r'^\|[\s\-:|]+\|$', l.strip())]
            
            if data_lines:
                rows = []
                for tl in data_lines:
                    cells = [c.strip() for c in tl.split('|')[1:-1]]
                    rows.append(cells)
                
                if rows:
                    num_cols = len(rows[0])
                    table = doc.add_table(rows=len(rows), cols=num_cols)
                    table.style = 'Table Grid'
                    
                    for ri, row in enumerate(rows):
                        for ci, cell in enumerate(row):
                            if ci < num_cols:
                                table.rows[ri].cells[ci].text = cell
                                for para in table.rows[ri].cells[ci].paragraphs:
                                    for run in para.runs:
                                        set_chinese_font(run, '宋体', 9)
            
            i = j
            continue
        
        # Handle regular paragraphs
        elif line.strip():
            flush_pending()
            text = line.strip()
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            
            p = doc.add_paragraph()
            run = p.add_run(text)
            set_chinese_font(run, '宋体', 10.5)
            p.paragraph_format.first_line_indent = Pt(21)
            p.paragraph_format.line_spacing = 1.5
        
        i += 1
    
    flush_pending()
    doc.save(output_file)
    print(f"Successfully converted to: {output_file}")
    shutil.rmtree(temp_dir, ignore_errors=True)
    return output_file


if __name__ == "__main__":
    # Default paths for this repository
    script_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(script_dir, "软件设计说明书.md")
    output_file = os.path.join(script_dir, "软件设计说明书.docx")
    mermaid_images_dir = os.path.join(script_dir, "figs", "md_figs")
    
    convert_md_to_word(
        md_file, 
        output_file, 
        title="无人水面艇仿真与强化学习训练系统（USVSIM）V1.0",
        base_dir=script_dir,
        mermaid_dir=mermaid_images_dir  # Use pre-rendered mermaid diagram images
    )